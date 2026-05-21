import os
from model.car import Carro
from service.car_service import Carservice
from model.aluguel import Aluguel
from run import app
from flask import request, jsonify
import pandas as pd
from flask import send_file
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList

UPLOAD_FOLDER = "uploads"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)  


@app.route('/car', methods=['POST'])
def post_carro():
    marca = request.form.get("marca")
    modelo = request.form.get("modelo")
    ano = int(request.form.get("ano") or 0)
    placa = request.form.get("placa")
    cor = request.form.get("cor")
    categoria = request.form.get("categoria")
    cambio = request.form.get("cambio")
    combustivel = request.form.get("combustivel")
    itens = request.form.get("itens") or ""
    preco_por_dia = float(request.form.get("preco_por_dia") or 0.0)

    imagens = request.files.getlist("foto")
    caminhos_imagens = []
    for img in imagens:
        if img.filename != "":
            caminho = os.path.join(UPLOAD_FOLDER, img.filename)
            img.save(caminho)
            caminhos_imagens.append(caminho)

    carro = Carservice.cadastra_carro(
        marca, modelo, ano, placa, cor, categoria, cambio, combustivel, itens, caminhos_imagens, preco_por_dia
    )

    if carro:
        return jsonify({"message": "Carro cadastrado"}), 201

    return jsonify({"message": "Erro ao cadastrar carro"}), 400


@app.route('/allcars', methods=['GET'])
def get_cars():
    carros = Carservice.listar_carros()
    if carros:
        return jsonify([car.to_dict_car() for car in carros]), 200
    return jsonify({"message": "Carros não encontrados"}), 404



@app.route("/alugar", methods=["POST"])
def alugar_carro():
    data = request.get_json()
    usuario_id = data.get("usuario_id")
    carro_id = data.get("carro_id")
    dias = data.get("dias", 1)  # padrão 1 dia se não informado

    if not usuario_id or not carro_id:
        return jsonify({"erro": "usuario_id e carro_id são obrigatórios"}), 400

    resposta, status = Carservice.alugar_carro(usuario_id, carro_id, dias)
    return jsonify(resposta), status


@app.route("/meusalugueis", methods=["GET"])
def meus_alugueis():
    usuario_id = request.args.get("usuario_id")

    if not usuario_id:
        return jsonify({"erro": "É necessário informar o usuarioID na URL"}), 400
    
    alugueis = Aluguel.query.filter_by(usuario_id=usuario_id, status="ativo").all()

    # alugueis = Aluguel.query.filter_by(usuario_id=usuario_id).all()
    # if not alugueis:
    #     return jsonify([]), 200

    resultado = []
    for aluguel in alugueis:
        carro = Carro.query.get(aluguel.carro_id)
        if carro:
            carro_dict = carro.to_dict_car()
            carro_dict["aluguel_id"] = aluguel.id
            carro_dict["data_inicio"] = aluguel.data_inicio.strftime("%Y-%m-%d") if aluguel.data_inicio else None
            carro_dict["data_fim"] = aluguel.data_fim.strftime("%Y-%m-%d") if aluguel.data_fim else None
            carro_dict["status"] = aluguel.status
            carro_dict["valor_total"] = aluguel.valor_total
            resultado.append(carro_dict)

    return jsonify(resultado), 200

@app.route("/devolver/<int:aluguel_id>", methods=["PUT"])
def devolver_carro(aluguel_id):
    resultado, status = Carservice.devolver_carro(aluguel_id)
    return jsonify(resultado), status

@app.route("/allalugueis", methods=["GET"])
def all_alugueis():
    alugueis = Aluguel.query.all()

    resultado = []
    for aluguel in alugueis:
        resultado.append({
            "id": aluguel.id,
            "usuario_id": aluguel.usuario_id,
            "carro_id": aluguel.carro_id,
            "status": aluguel.status,
            "data_inicio": aluguel.data_inicio.strftime("%Y-%m-%d") if aluguel.data_inicio else None,
            "data_fim": aluguel.data_fim.strftime("%Y-%m-%d") if aluguel.data_fim else None,
            "valor_total": aluguel.valor_total
        })

    return jsonify(resultado), 200

@app.route("/car/<int:carro_id>", methods=["DELETE"])
def deletar_carro(carro_id):
    resposta, status = Carservice.deletar_carro(carro_id)
    return jsonify(resposta), status

@app.route("/carros", methods=["GET"])
def buscar_carros():

    termo = request.args.get("busca")
    ano = request.args.get("ano")
    ordem = request.args.get("ordem")

    carros = Carservice.pesquisar_carros(
        termo=termo,
        ano=ano,
        ordem=ordem
    )

    return jsonify([car.to_dict_car() for car in carros]), 200

@app.route("/avaliar", methods=["POST"])
def avaliar_carro():

    data = request.get_json()

    usuario_id = data.get("usuario_id")
    carro_id = data.get("carro_id")
    nota = data.get("nota")
    comentario = data.get("comentario")

    if not usuario_id or not carro_id or not nota:
        return jsonify({"erro": "usuario_id, carro_id e nota são obrigatórios"}), 400

    resposta, status = Carservice.avaliar_carro(
        usuario_id,
        carro_id,
        int(nota),
        comentario
    )

    return jsonify(resposta), status

@app.route("/carros/<int:carro_id>/media", methods=["GET"])
def obter_media(carro_id):

    carro = Carro.query.get(carro_id)

    if not carro:
        return jsonify({"erro": "Carro não encontrado"}), 404

    return jsonify({
        "media": carro.media_avaliacao or 0,
        "total": carro.total_avaliacoes or 0
    }), 200

@app.route("/exportar-excel", methods=["GET"])
def exportar_excel():

    

    carros = Carro.query.all()
    alugueis = Aluguel.query.all()

    # =========================
    # DADOS DOS CARROS
    # =========================
    dados_carros = []

    for c in carros:
        dados_carros.append({
            "Marca": c.marca,
            "Modelo": c.modelo,
            "Ano": c.ano,
            "Preço por dia": c.preco_por_dia,
            "Disponível": "Sim" if c.disponivel else "Não"
        })

    df_carros = pd.DataFrame(dados_carros)

    # =========================
    # DADOS DOS ALUGUÉIS
    # =========================
    dados_alugueis = []

    for a in alugueis:
        dados_alugueis.append({
            "Carro ID": a.carro_id,
            "Usuário ID": a.usuario_id,
            "Status": a.status,
            "Data Início": str(a.data_inicio),
            "Data Fim": str(a.data_fim),
            "Valor Total": a.valor_total
        })

    df_alugueis = pd.DataFrame(dados_alugueis)

    # =========================
    # DADOS DO GRÁFICO
    # =========================
    disponiveis = len([c for c in carros if c.disponivel])
    alugados = len([c for c in carros if not c.disponivel])

    # =========================
    # CRIAR EXCEL
    # =========================
    caminho = "relatorio_alugae.xlsx"

    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:

        # =========================
        # ABA 1 - CARROS
        # =========================
        df_carros.to_excel(writer, sheet_name="Carros", index=False)

        # =========================
        # ABA 2 - ALUGUÉIS
        # =========================
        df_alugueis.to_excel(writer, sheet_name="Alugueis", index=False)

        # =========================
        # ABA 3 - GRÁFICOS
        # =========================
        wb = writer.book

        ws_grafico = wb.create_sheet("Graficos")

        # Título
        ws_grafico.merge_cells("A1:B1")
        ws_grafico["A1"] = "Dashboard do Sistema"

        # Cabeçalhos
        ws_grafico.append(["Status", "Quantidade"])

        # Dados
        ws_grafico.append(["Disponíveis", disponiveis])
        ws_grafico.append(["Alugados", alugados])

        # =========================
        # CRIAR GRÁFICO PIZZA
        # =========================
        grafico = PieChart()

        grafico.title = "Status dos Carros"

        labels = Reference(
            ws_grafico,
            min_col=1,
            min_row=3,
            max_row=4
        )

        data = Reference(
            ws_grafico,
            min_col=2,
            min_row=2,
            max_row=4
        )

        grafico.add_data(data, titles_from_data=True)
        grafico.set_categories(labels)

        # Mostrar porcentagem
        grafico.dataLabels = DataLabelList()
        grafico.dataLabels.showPercent = True
        grafico.dataLabels.showVal = False

        # Tamanho
        grafico.width = 12
        grafico.height = 8

        # Adiciona gráfico
        ws_grafico.add_chart(grafico, "D3")

        # =========================
        # AJUSTAR LARGURA AUTOMÁTICA
        # =========================
        for sheet_name in writer.sheets:

            sheet = writer.sheets[sheet_name]

            for col in sheet.columns:

                max_length = 0
                first_cell = None

                for cell in col:
                    if hasattr(cell, "column_letter"):
                        first_cell = cell
                        break

                if not first_cell:
                    continue

                col_letter = first_cell.column_letter

                for cell in col:
                    try:
                        if cell.value:
                            max_length = max(
                                max_length,
                                len(str(cell.value))
                            )
                    except:
                        pass

                sheet.column_dimensions[col_letter].width = max_length + 4

    return send_file(caminho, as_attachment=True)